"""FC --selftest: the release check the build workflow runs on every system FC supports.

It loads every bundled library, runs each kind of tool on small sample files made on the
spot (images, PDF, Word export, PowerPoint, ZIP, audio, video), then opens the octagon,
walks a menu and closes it again. One line per check goes to the console and to
selftest.txt in FC's log folder (a windowed exe has no console); the exit code is the
number of failed checks. --no-gui skips the window.
"""
from __future__ import annotations

import importlib
import os
import platform
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path

import fc

LINES: list[str] = []
FAILED: list[str] = []


def say(line: str):
    LINES.append(line)
    try:
        print(line, flush=True)
    except (AttributeError, OSError):          # windowed exe: no console to print to
        pass


def check(name: str, fn):
    started = time.monotonic()
    try:
        detail = fn()
    except Exception as exc:  # noqa: BLE001 - every failure is reported, the rest still run
        FAILED.append(name)
        say(f"FAIL  {name}: {type(exc).__name__}: {exc}")
        for line in traceback.format_exc().strip().splitlines()[-4:]:
            say(f"        {line}")
        return None
    say(f"ok    {name} ({time.monotonic() - started:.1f} s){f'  {detail}' if detail else ''}")
    return detail


def _one(paths) -> Path:
    """The single result a tool returned, which must exist and not be empty."""
    assert len(paths) == 1, f"expected one result, got {len(paths)}"
    out = Path(paths[0])
    assert out.exists(), f"{out.name} was not written"
    if out.is_file():
        assert out.stat().st_size > 0, f"{out.name} is empty"
    else:
        assert any(out.iterdir()), f"{out.name} is an empty folder"
    return out


def _smaller_or_already_small(fn):
    """Compress tools refuse to hand back a bigger file; that answer is fine too."""
    try:
        return _one(fn()).name
    except ValueError as exc:
        if "small" not in str(exc):
            raise
        return "already small"


# ------------------------------------------------------------------ the tools
def _tools(tmp: Path):
    from PIL import Image, ImageDraw

    import converters
    import file_tools
    import history

    run = file_tools.run

    def sample_image():
        img = Image.new("RGB", (900, 600))
        draw = ImageDraw.Draw(img)
        for x in range(0, 900, 6):                              # a busy picture compresses well
            draw.line([(x, 0), (900 - x, 600)], fill=(x % 256, (x * 3) % 256, (x * 7) % 256), width=3)
        draw.text((40, 40), "FC self-test", font=converters.font(48, bold=True), fill="white")
        img.save(tmp / "sample.png")
        return "sample.png"

    check("make a sample image", sample_image)
    png = tmp / "sample.png"
    for ext in ("jpg", "webp", "pdf", "ico", "bmp"):
        check(f"image -> {ext}", lambda ext=ext: _one(run("fmt:" + ext, [png])).name)
    check("image watermark (font)", lambda: _one(run("img_watermark", [png], "TEST")).name)
    check("image resize / rotate / split",
          lambda: [_one(run(k, [png])).name for k in ("img_resize", "img_rotate", "img_split")])
    check("image compress", lambda: _smaller_or_already_small(lambda: run("img_compress", [png])))

    def sample_pdf():
        import pymupdf

        for name in ("a.pdf", "b.pdf"):
            doc = pymupdf.open()
            for i in range(3):
                page = doc.new_page()
                page.insert_text((72, 90), f"FC self-test {name} page {i + 1}", fontsize=18)
                page.insert_image(pymupdf.Rect(72, 120, 400, 340), filename=str(png))
            doc.save(tmp / name)
            doc.close()
        return "a.pdf, b.pdf"

    check("make sample PDFs", sample_pdf)
    a, b = tmp / "a.pdf", tmp / "b.pdf"
    check("pdf merge", lambda: _one(run("merge", [a, b])).name)
    check("pdf split", lambda: f"{len(list(_one(run('split', [a])).iterdir()))} pages")
    check("pdf extract / remove pages",
          lambda: [_one(run(k, [a], "2")).name for k in ("extract", "remove")])
    check("pdf rotate / reverse / repair",
          lambda: [_one(run(k, [a])).name for k in ("rot_r", "reverse", "repair")])
    check("pdf watermark / page numbers",
          lambda: [_one(run(k, [a], "TEST")).name for k in ("watermark", "numbers")])

    def protect_unlock():
        locked = _one(run("protect", [a], "pass1234"))
        return _one(run("unlock", [locked], "pass1234")).name

    check("pdf protect + unlock", protect_unlock)
    check("pdf strip metadata", lambda: _one(run("strip", [a])).name)
    check("pdf compress", lambda: _smaller_or_already_small(lambda: run("compress", [a])))

    def pdf_text():
        out = _one(run("to_txt", [a]))
        assert "page 2" in out.read_text(encoding="utf-8"), "the text did not come out"
        return out.name

    check("pdf -> text", pdf_text)
    check("pdf -> png pages", lambda: _one(run("to_png", [a])).name)
    check("pdf -> Word (pdf2docx)", lambda: _one(run("to_word", [a])).name)
    check("pdf -> PowerPoint", lambda: _one(run("to_pptx", [a])).name)
    check("pdf get images", lambda: _one(run("images", [a])).name)
    check("images -> one PDF", lambda: _one(run("img_pdf", [png, tmp / "sample.jpg"])).name)

    def office_compress():
        from pptx import Presentation

        big = Image.effect_noise((1400, 1400), 80).convert("RGB")
        big.save(tmp / "noise.png")
        pres = Presentation()
        pres.slides.add_slide(pres.slide_layouts[6]).shapes.add_picture(str(tmp / "noise.png"), 0, 0)
        pres.save(tmp / "deck.pptx")
        return _smaller_or_already_small(lambda: run("office_compress", [tmp / "deck.pptx"]))

    check("PowerPoint compress (repack)", office_compress)
    check("zip", lambda: _one(run("zip", [png, a])).name)

    def zip_lock():
        import pyzipper

        out = _one(run("zip_lock", [png], "pass1234"))
        with pyzipper.AESZipFile(out) as z:
            z.setpassword(b"pass1234")
            assert z.read("sample.png") == png.read_bytes(), "the locked ZIP does not open"
        return out.name

    check("zip with password (AES)", zip_lock)

    def ffmpeg_samples():
        import imageio_ffmpeg

        exe = imageio_ffmpeg.get_ffmpeg_exe()
        for args in (["-f", "lavfi", "-i", "sine=frequency=440:duration=3", str(tmp / "tone.wav")],
                     ["-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=15", "-f", "lavfi",
                      "-i", "sine=duration=2", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p",
                      "-c:a", "aac", str(tmp / "clip.mp4")]):
            subprocess.run([exe, "-hide_banner", "-loglevel", "error", "-y", *args], check=True,
                           timeout=120, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        return Path(exe).name

    check("ffmpeg makes sample audio and video", ffmpeg_samples)
    wav, mp4 = tmp / "tone.wav", tmp / "clip.mp4"
    for ext in ("mp3", "flac", "ogg", "m4a"):
        check(f"audio -> {ext}", lambda ext=ext: _one(run("fmt:" + ext, [wav])).name)
    for ext in ("webm", "mkv", "gif", "mp3"):
        check(f"video -> {ext}", lambda ext=ext: _one(run("fmt:" + ext, [mp4])).name)
    check("audio compress", lambda: _smaller_or_already_small(lambda: run("media_compress", [wav])))
    check("media merge", lambda: _one(run("media_merge", [mp4, tmp / "clip.mkv"])).name)
    check("before / after card", lambda: history.compare(png, tmp / "sample.jpg").name)

    if not fc.IS_WIN:                     # optional helpers on Linux: tested only where installed
        import shutil

        if shutil.which("soffice") or shutil.which("libreoffice"):
            check("LibreOffice: PowerPoint -> PDF", lambda: _one(run("fmt:pdf", [tmp / "deck.pptx"])).name)
        else:
            say("skip  LibreOffice is not installed (Office conversions on Linux need it)")


# ------------------------------------------------------------------ the window
def _gui():
    from PyQt6.QtCore import QTimer
    from PyQt6.QtQuick import QQuickView
    from PyQt6.QtWidgets import QApplication

    if not fc.IS_WIN and not (os.environ.get("DISPLAY") or os.environ.get("WAYLAND_DISPLAY")):
        say("skip  window: no display (run under X11, Wayland or Xvfb to test it)")
        return
    qapp = QApplication.instance() or QApplication(sys.argv[:1] + fc.qt_platform())
    say(f"info  Qt platform: {qapp.platformName()}")
    state: dict = {}

    def build():
        state["app"] = fc.App(qapp)
        return ""

    if check("open the octagon", build) is None:
        return
    app = state["app"]
    view = app.hexagon.view

    def wait(seconds: float, until=lambda: False):
        end = time.monotonic() + seconds
        while time.monotonic() < end and not until():
            qapp.processEvents()
            time.sleep(0.01)

    wait(15, lambda: app.hexagon.shown_once > 0)

    def shown():
        assert view.status() == QQuickView.Status.Ready, f"QML status {view.status()}"
        errors = [e.toString() for e in view.errors()]
        assert not errors, "; ".join(errors)
        assert app.hexagon.shown_once > 0, "the octagon never reached the screen"
        image = view.grabWindow()
        assert not image.isNull() and image.width() > 0, "the window has no picture"
        return f"on screen {app.hexagon.shown_once:.2f} s after FC started"

    check("octagon on screen", shown)

    def menus():
        backend = app.hexagon.backend
        labels = [o["label"] for o in backend.options]
        assert labels == [label for label, _ in fc.HOME], labels
        app.open_for_files([Path(tempfile.gettempdir()) / "x.pdf"])     # the PDF tools
        assert backend.title == "PDF" and len(backend.options) == fc.SIDES, backend.title
        backend.pick([o["label"] for o in backend.options].index("Convert"))
        assert "Word" in [o["label"] for o in backend.options], "the Convert group did not open"
        backend.back()
        backend.back()
        assert backend.atHome, "Back did not return to the start"
        wait(0.5)
        return "home -> PDF -> Convert -> back"

    check("menus", menus)
    QTimer.singleShot(0, app.hexagon.quit)
    wait(3)


def run(gui: bool = True) -> int:
    say(f"FC {fc.VERSION} self-test on {platform.platform()} (Python {platform.python_version()})")
    if not fc.IS_WIN:
        say(f"info  glibc {' '.join(platform.libc_ver())}; session "
            f"{os.environ.get('XDG_SESSION_TYPE', '?')}; desktop {os.environ.get('XDG_CURRENT_DESKTOP', '?')}")
    modules = ["PyQt6.QtQuick", "PyQt6.QtQml", "PyQt6.QtNetwork", "PIL.Image", "pymupdf", "pyzipper",
               "pptx", "pdf2docx", "cv2", "numpy", "imageio_ffmpeg",
               "comtypes.client" if fc.IS_WIN else "PyQt6.QtDBus"]
    for name in modules:
        check(f"import {name}", lambda name=name: importlib.import_module(name) and "")
    def bundled():
        missing = [n for n in ("hex.qml", "fc.png", "fc-mark.png", "icons/repeat.svg") if not (fc.HERE / n).exists()]
        assert not missing, f"missing {', '.join(missing)}"
        return ""

    check("bundled files", bundled)
    with tempfile.TemporaryDirectory(prefix="fc-selftest-") as tmp:
        _tools(Path(tmp))
    if gui:
        _gui()
    say(f"FC {fc.VERSION} selftest {'ok' if not FAILED else 'FAILED: ' + ', '.join(FAILED)}")
    try:
        fc.LOG_DIR.mkdir(parents=True, exist_ok=True)
        (fc.LOG_DIR / "selftest.txt").write_text("\n".join(LINES) + "\n", encoding="utf-8")
    except OSError:
        pass
    return min(len(FAILED), 100)
